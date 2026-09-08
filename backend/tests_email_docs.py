#!/usr/bin/env python3
"""
Email documents: extraction, safe filing paths, and the update whitelist.

The whitelist is the part with teeth: a founder's document is excellent
evidence for some fields and terrible evidence for others, and a wrong
overwrite corrupts a verified record. Every doubtful case must be a no-op.
"""
import email
import email.mime.application
import email.mime.image
import email.mime.multipart
import email.mime.text
import os
import sys
import warnings

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("GCP_PROJECT_ID", "averroes-deal-origination")

from services.email_docs_service import (  # noqa: E402
    doc_gcs_path, extract_attachments, sanitize_filename,
)

fails = 0


def chk(label, got, want):
    global fails
    ok = got == want
    print(("PASS" if ok else "FAIL"), label, "->", got, "" if ok else f"(wanted {want})")
    if not ok:
        fails += 1


print("── Filenames are made path-safe ──")
chk("plain name survives", sanitize_filename("Deck 2026.pdf"), "Deck 2026.pdf")
# Traversal collapses to the LAST segment - the only part that is a filename.
chk("path traversal is stripped", sanitize_filename("../../etc/passwd"), "passwd")
chk("separators are stripped", sanitize_filename("a/b\\c.pdf"), "c.pdf")
chk("empty becomes a name", sanitize_filename(""), "attachment")
chk("bounded length", len(sanitize_filename("x" * 300)), 120)

print()
print("── GCS paths are stable and browsable ──")
chk("company/date/filename shape",
    doc_gcs_path("Acme Ltd", "2026-08-20T10:00:00+00:00", "deck.pdf"),
    "email-docs/Acme Ltd/2026-08-20_deck.pdf")
chk("no date -> undated, never a crash",
    doc_gcs_path("Acme", "", "a.pdf"), "email-docs/Acme/undated_a.pdf")
chk("hostile company names are neutralised",
    "/" not in doc_gcs_path("A/B../C", "2026-01-01", "x.pdf").replace("email-docs/", "", 1).split("/")[0], True)

print()
print("── Attachment extraction from a real MIME message ──")
msg = email.mime.multipart.MIMEMultipart()
msg.attach(email.mime.text.MIMEText("Hi, deck attached."))
pdf = email.mime.application.MIMEApplication(b"%PDF-fake", _subtype="pdf")
pdf.add_header("Content-Disposition", "attachment", filename="Acme Deck.pdf")
msg.attach(pdf)
logo = email.mime.image.MIMEImage(b"\x89PNG-fake", _subtype="png")
logo.add_header("Content-Disposition", "inline", filename="logo.png")
msg.attach(logo)
parsed = email.message_from_bytes(msg.as_bytes())
atts = extract_attachments(parsed)
chk("both files found (everything attached, per Ishu)", len(atts), 2)
chk("pdf named and typed", (atts[0]["filename"], atts[0]["content_type"]),
    ("Acme Deck.pdf", "application/pdf"))
chk("bytes preserved", atts[0]["data"], b"%PDF-fake")
chk("a body-only email yields nothing",
    extract_attachments(email.message_from_string("Subject: hi\n\njust text")), [])

print()
print("── Cost guards: what earns an AI read ──")
from services.email_docs_service import AI_READS_PER_RUN, MIN_AI_IMAGE_BYTES, should_analyse  # noqa: E402
chk("a PDF is always read (decks and accounts live there)",
    should_analyse("application/pdf", 5_000), True)
chk("a signature-logo-sized image is filed but NOT read",
    should_analyse("image/png", 30 * 1024), False)
chk("a large image (scanned doc, chart) IS read",
    should_analyse("image/png", 400 * 1024), True)
chk("office docs ARE read now (converted to text first, Document SmartFill)",
    should_analyse("application/vnd.openxmlformats-officedocument.wordprocessingml.document", 900_000), True)
chk("a random binary type is still filed only",
    should_analyse("application/zip", 900_000, "archive.zip"), False)
chk("the image threshold is sane (50-500KB)",
    50 * 1024 <= MIN_AI_IMAGE_BYTES <= 500 * 1024, True)
chk("the per-run read budget is bounded", 1 <= AI_READS_PER_RUN <= 50, True)

print()
print("── Document SmartFill: fill vs confirm (services/doc_smartfill.py) ──")
from services.doc_smartfill import (  # noqa: E402
    COLUMN_TYPES, office_kind, office_to_text, plan_updates, merge_writes,
)
company = {"name": "Acme", "revenue_estimate_m": 3.0, "employees": 20,
           "description": "A B2B SaaS platform for logistics teams.",
           "sector": "Logistics Tech", "website": "", "hq_city": None,
           "revenue_y1": 4_000_000, "revenue_y1_date": "2024-03-31",
           "revenue_y2": 3_000_000, "revenue_y2_date": "2023-03-31"}

ex = {"summary": "FY25 management accounts.",
      "company": {"website": "https://acme.co.uk", "employees": "34", "sector": "Fintech",
                  "description": "SaaS company.", "year_founded": 2016, "total_raised_m": 4.5,
                  "ebitda_margin_pct": -12.5},
      "evidence": {"website": "footer p1", "employees": "34 FTEs, slide 3", "sector": "slide 2"},
      "financial_years": [
          {"period_end": "2025-03-31", "basis": "actual", "revenue": 5_200_000, "gross_profit": 4_160_000,
           "profit_before_tax": -250_000, "cash": 812_345, "evidence": "P&L p4"},
          {"period_end": "2026-03-31", "basis": "forecast", "revenue": 9_000_000},
      ]}
plan = plan_updates(company, ex)
fills = {i["key"]: i for i in plan["fills"]}
conf = {i["key"]: i for i in plan["conflicts"]}

chk("blank website is FILLED automatically", fills["website"]["writes"], {"website": "https://acme.co.uk"})
chk("blank year_founded / total_raised / margin are fills",
    {"year_founded", "total_raised_m", "ebitda_margin_pct"} <= set(fills), True)
chk("negative EBITDA margin allowed (signed field)", fills["ebitda_margin_pct"]["writes"]["ebitda_margin_pct"], -12.5)
chk("existing employees 20 vs document 34 is a CONFLICT, not a write", conf["employees"]["writes"], {"employees": 34})
chk("conflict shows current and document values", (conf["employees"]["old"], conf["employees"]["new"]), ("20", "34"))
chk("conflict carries the evidence", conf["employees"]["evidence"], "34 FTEs, slide 3")
chk("existing sector vs different sector is a conflict", "sector" in conf, True)
chk("a SHORTER description never wins and is not even a conflict",
    "description" in fills or "description" in conf, False)
chk("a longer description is a fill (longer-wins doctrine)",
    "description" in {i["key"] for i in plan_updates(company, {"company": {
        "description": company["description"] + " Serves 200 enterprise customers across the UK."}})["fills"]}, True)

# ── Multi-year store: cells, not slots ──
from services.doc_smartfill import cells_from_columns, project_to_columns  # noqa: E402
fin_fill = fills.get("financials:2025-03-31:new")
chk("a year we do not hold is a FILL item, grouped per period", fin_fill is not None, True)
cells = fin_fill["writes"]["financials"]
chk("all four stated metrics become cells",
    sorted(c["metric"] for c in cells), ["cash", "gross_profit", "profit_before_tax", "revenue"])
chk("loss kept as a negative PBT cell", next(c["value"] for c in cells if c["metric"] == "profit_before_tax"), -250_000.0)
chk("forecast year is a separate, flagged fill (never actual)",
    fills["financials:2026-03-31:new"]["writes"]["financials"][0]["basis"], "forecast")
chk("...and its label says so", "(forecast)" in fills["financials:2026-03-31:new"]["label"], True)
chk("derived latest revenue conflicts with the stored 3.0m estimate",
    conf["revenue_estimate_m"]["writes"]["revenue_estimate_m"], 5.2)
chk("derived growth uses store + new year: 4.0m -> 5.2m = +30% (fill, nothing stored)",
    fills["revenue_growth_pct"]["writes"]["revenue_growth_pct"], 30.0)

# Same year, same number -> nothing; different number -> conflict cell.
same = plan_updates(company, {"financial_years": [{"period_end": "2024-03-31", "basis": "actual", "revenue": 4_010_000}]})
chk("a held cell within 1% is NOT a change", same, {"fills": [], "conflicts": []})
diff = plan_updates(company, {"financial_years": [{"period_end": "2024-03-31", "basis": "actual", "revenue": 4_800_000}]})
chk("a held cell that DISAGREES is a conflict for that year", [i["key"] for i in diff["conflicts"]][0], "financials:2024-03-31:changed")
chk("...figures derived from an unconfirmed year wait too", diff["fills"], [])
chk("conflict shows stored vs document with the stored source",
    "£4.00m" in diff["conflicts"][0]["old"] and "£4.80m" in diff["conflicts"][0]["new"], True)
# A year we hold gains a metric we lacked: fill, no conflict.
gain = plan_updates(company, {"financial_years": [{"period_end": "2024-03-31", "basis": "actual",
                                                    "revenue": 4_000_000, "gross_margin_pct": 78.5}]})
chk("new metric on a held year is a fill; the matching revenue is silent",
    ([i["key"] for i in gain["fills"] if i["key"].startswith("financials")],
     [c["metric"] for c in gain["fills"][0]["writes"]["financials"]]),
    (["financials:2024-03-31:new"], ["gross_margin_pct"]))
# Segments ride as revenue cells with a segment name.
seg = plan_updates({"name": "X"}, {"financial_years": [{"period_end": "2025-12-31", "basis": "actual", "revenue": 3_200_000,
                                                          "segments": [{"name": "Instruments", "revenue": 2_000_000},
                                                                       {"name": "Software", "revenue": 1_200_000}]}]})
sc = seg["fills"][0]["writes"]["financials"]
chk("segment split stored as revenue cells with segment", sorted(c["segment"] for c in sc), ["", "Instruments", "Software"])
# When the store is loaded (company['_financials']) it is the truth, not the columns.
stored = dict(company, _financials=[{"period_end": "2024-03-31", "metric": "revenue", "segment": "", "value": 4_000_000,
                                     "unit": "GBP", "basis": "actual", "source": "Companies House"}])
chk("store rows take precedence over legacy columns",
    plan_updates(stored, {"financial_years": [{"period_end": "2023-03-31", "basis": "actual", "revenue": 3_000_000}]})["fills"][0]["key"],
    "financials:2023-03-31:new")

print()
print("── Legacy columns <-> store: seed and projection ──")
seeded = cells_from_columns({"revenue_y1": 4_000_000, "revenue_y1_date": "2024-03-31", "cash_y1": 500_000,
                             "revenue_y2": 3_000_000, "revenue_y2_date": "2023-03-31", "employees_ch": 28}, "Companies House")
chk("seed from columns yields dated cells", sorted((c["period_end"], c["metric"]) for c in seeded),
    [("2023-03-31", "revenue"), ("2024-03-31", "cash"), ("2024-03-31", "employees"), ("2024-03-31", "revenue")])
chk("undated slots are not seeded", cells_from_columns({"revenue_y1": 1}, "x"), [])
proj = project_to_columns([
    {"period_end": "2025-12-31", "metric": "revenue", "segment": "", "value": 3_200_000, "basis": "actual"},
    {"period_end": "2025-12-31", "metric": "ebitda", "segment": "", "value": -320_000, "basis": "actual"},
    {"period_end": "2024-12-31", "metric": "revenue", "segment": "", "value": 1_700_000, "basis": "actual"},
    {"period_end": "2023-12-31", "metric": "revenue", "segment": "", "value": 900_000, "basis": "actual"},
    {"period_end": "2022-12-31", "metric": "revenue", "segment": "", "value": 500_000, "basis": "actual"},
    {"period_end": "2026-12-31", "metric": "revenue", "segment": "", "value": 6_000_000, "basis": "forecast"},
    {"period_end": "2025-12-31", "metric": "revenue", "segment": "Software", "value": 1_000_000, "basis": "actual"},
])
chk("projection = latest three ACTUAL whole-company years",
    (proj["revenue_y1"], proj["revenue_y1_date"], proj["revenue_y2"], proj["revenue_y3"], proj["revenue_y3_date"]),
    (3_200_000.0, "2025-12-31", 1_700_000.0, 900_000.0, "2023-12-31"))
chk("forecast and segment rows never reach the legacy columns", 6_000_000.0 not in proj.values() and 1_000_000.0 not in proj.values(), True)
chk("EBITDA margin derived for the latest year", proj["ebitda_margin_pct"], -10.0)
chk("empty store clears the slots (None), never leaves stale values", project_to_columns([])["revenue_y1"], None)

# Refinements from the first live review (Plastometrex, 8 Sep 2026)
sub = plan_updates({"name": "X", "active_investors": "EMV Capital, Innovate UK, Vanneck",
                    "website": "https://plastometrex.com"},
                   {"company": {"active_investors": {"value": "Innovate UK", "evidence": "p2"},
                                "website": {"value": "www.Plastometrex.com", "evidence": "footer"}}})
chk("a SUBSET of stored investors is nothing (no conflict, no fill)", sub, {"fills": [], "conflicts": []})
add = plan_updates({"name": "X", "active_investors": "EMV Capital, Innovate UK"},
                   {"company": {"active_investors": {"value": "Innovate UK, Martlet Capital", "evidence": "p2"}}})
chk("NEW investors are merged in as a fill, nothing removed",
    (add["conflicts"], add["fills"][0]["writes"]), ([], {"active_investors": "EMV Capital, Innovate UK, Martlet Capital"}))
chk("...and the label names what was added", "added: Martlet Capital" in add["fills"][0]["label"], True)
chk("{value, evidence} shape carries the evidence through",
    plan_updates({"name": "X", "employees": 20}, {"company": {"employees": {"value": 34, "evidence": "slide 3"}}})["conflicts"][0]["evidence"], "slide 3")
cagr = plan_updates({"name": "X"}, {"company": {"revenue_cagr_3yr_pct": {"value": 92, "evidence": "p6"}}})
chk("multi-year CAGR lands in its own column, not last-year growth",
    cagr["fills"][0]["writes"], {"revenue_cagr_3yr_pct": 92.0})
chk("fields off the schema are ignored entirely",
    plan_updates(company, {"company": {"status": "Won", "averroes_fit_score": 1.0, "contact_email": "x@y.z"}}),
    {"fills": [], "conflicts": []})
chk("zero/negative headcount and garbage numbers are refused",
    plan_updates({"name": "X"}, {"company": {"employees": -5, "total_raised_m": "about forty"}}),
    {"fills": [], "conflicts": []})
chk("a no-change proposal is dropped", plan_updates(company, {"company": {"employees": 20}}), {"fills": [], "conflicts": []})
chk("empty extraction is safe", plan_updates(company, {}), {"fills": [], "conflicts": []})
chk("None extraction is safe", plan_updates(company, None), {"fills": [], "conflicts": []})
chk("every write column has a declared BigQuery type (or is the virtual financials list)",
    all(c in COLUMN_TYPES or c == "financials" for i in plan["fills"] + plan["conflicts"] for c in i["writes"]), True)
chk("merge_writes flattens items", merge_writes([{"writes": {"a": 1}}, {"writes": {"b": 2}}]), {"a": 1, "b": 2})

print()
print("── Office files become text ──")
chk("pptx by content type", office_kind("application/vnd.openxmlformats-officedocument.presentationml.presentation", "x"), "pptx")
chk("octet-stream + .xlsx falls back to the extension", office_kind("application/octet-stream", "Model v3.XLSX"), "xlsx")
chk("pdf is not an office kind", office_kind("application/pdf", "deck.pdf"), None)
import io as _io
from pptx import Presentation as _P
_prs = _P(); _s = _prs.slides.add_slide(_prs.slide_layouts[5]); _s.shapes.title.text = "FY25 revenue GBP 5.2m"
_b = _io.BytesIO(); _prs.save(_b)
chk("pptx text extracted with slide labels", "Slide 1" in office_to_text("pptx", _b.getvalue())
    and "5.2m" in office_to_text("pptx", _b.getvalue()), True)
from openpyxl import Workbook as _W
_wb = _W(); _ws = _wb.active; _ws.title = "P&L"; _ws.append(["Revenue", 5200000]); _ws.append(["EBITDA", -250000])
_b2 = _io.BytesIO(); _wb.save(_b2)
_t = office_to_text("xlsx", _b2.getvalue())
chk("xlsx rows extracted with sheet label", "Sheet P&L" in _t and "5200000" in _t, True)
chk("corrupt bytes never raise", office_to_text("docx", b"not a docx"), "")
chk("office files are worth an AI read", should_analyse("application/octet-stream", 10, "deck.pptx"), True)

print()
print("── PDF links in the body (the Plastometrex case) ──")
from services.email_docs_service import _is_safe_url, extract_pdf_links  # noqa: E402

chk("a direct pdf link is found",
    extract_pdf_links("Our deck: https://plastometrex.com/files/PIP-deck.pdf - enjoy"),
    ["https://plastometrex.com/files/PIP-deck.pdf"])
chk("query strings survive",
    extract_pdf_links("see https://a.co/x.pdf?dl=1&v=2 now"), ["https://a.co/x.pdf?dl=1&v=2"])
chk("drive/dropbox share links are NOT treated as pdfs (they serve HTML)",
    extract_pdf_links("https://drive.google.com/file/d/abc/view and https://www.dropbox.com/s/x/deck"), [])
chk("capped at 2 per email",
    len(extract_pdf_links(" ".join(f"https://a.co/{i}.pdf" for i in range(5)))), 2)
chk("duplicates collapse",
    extract_pdf_links("https://a.co/x.pdf and again https://a.co/x.pdf"), ["https://a.co/x.pdf"])
chk("no links, no crash", extract_pdf_links(""), [])

print()
print("── The SSRF guard: external mail must never reach anything internal ──")
pub = lambda h: ["93.184.216.34"]
chk("a normal public URL passes", _is_safe_url("https://plastometrex.com/deck.pdf", pub), True)
chk("http (not just https) is allowed", _is_safe_url("http://a.co/x.pdf", pub), True)
chk("the metadata server is refused",
    _is_safe_url("http://metadata.google.internal/computeMetadata/v1/x.pdf", pub), False)
chk("any .internal host is refused", _is_safe_url("https://x.svc.internal/a.pdf", pub), False)
chk("localhost is refused", _is_safe_url("http://localhost/x.pdf", pub), False)
chk("raw-IP hosts are refused outright", _is_safe_url("https://93.184.216.34/x.pdf", pub), False)
chk("a host resolving to the metadata IP is refused",
    _is_safe_url("https://evil.example/x.pdf", lambda h: ["169.254.169.254"]), False)
chk("a host resolving to a private range is refused",
    _is_safe_url("https://evil.example/x.pdf", lambda h: ["10.0.0.5"]), False)
chk("ONE private address among many public poisons the whole host",
    _is_safe_url("https://evil.example/x.pdf", lambda h: ["93.184.216.34", "127.0.0.1"]), False)
chk("credentials in the URL are refused", _is_safe_url("https://user@a.co/x.pdf", pub), False)
chk("non-http schemes are refused", _is_safe_url("ftp://a.co/x.pdf", pub), False)
chk("resolution failure fails CLOSED", _is_safe_url("https://a.co/x.pdf", lambda h: []), False)

print()
print(f"{fails} FAILURES" if fails else "ALL PASS")
sys.exit(1 if fails else 0)
